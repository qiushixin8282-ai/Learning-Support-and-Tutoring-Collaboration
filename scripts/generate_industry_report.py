from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "学习支持与辅导协作行业分析报告.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    for i, header in enumerate(headers):
        shade(tbl.rows[0].cells[i], "1F4E79")
        set_cell_text(tbl.rows[0].cells[i], header, True, "FFFFFF")
    for row_index, row in enumerate(rows):
        cells = tbl.add_row().cells
        for i, value in enumerate(row):
            if row_index % 2:
                shade(cells[i], "F3F6FA")
            set_cell_text(cells[i], value)
    if widths:
        for row in tbl.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return tbl


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("学习支持与辅导协作行业分析报告")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("面向在线教育学习服务的一期数智化转型建议").italic = True

    heading(doc, "结论摘要")
    paragraph(doc, "学习支持与辅导的关键难点不在于单次答复能力，而在于跨渠道请求的分拣、规则核对、跨角色交接和过程留痕。建议以受控任务协作为一期切入：先统一请求与责任状态，再将授权、规则版本和人工复核作为不可跳过的控制点。")
    paragraph(doc, "一期应在单课程、单地区、脱敏样例的受控范围内验证处理时长、返工次数、版本正确率和人工复核成本。没有完成样本基线和权限确认前，不应承诺具体业务收益。")

    heading(doc, "行业环境与业务变化")
    paragraph(doc, "在线教育的学习服务从课后答疑逐步延伸到学习提醒、进度跟踪、作业反馈、学习路径建议和特殊支持协同。服务对象和问题类型增多后，单靠个人经验和即时沟通难以稳定地保持规则一致、责任清晰和体验可追溯。")
    table(doc, ["变化", "对学习服务的影响", "一期应对"], [
        ["服务渠道增加", "请求分散在表单、即时通信和邮件，状态难以汇总", "统一入口和匿名任务编号"],
        ["规则频繁更新", "课程政策、服务口径和适用范围容易使用旧版本", "规则版本硬校验和变更留痕"],
        ["协作角色增多", "助教、运营、讲师和数据角色之间交接易遗漏", "具名责任人和异常复核队列"],
        ["合规要求提高", "学习信息需要遵循最小化、授权和审计原则", "最小字段、最小权限和人工控制"],
    ], [3.2, 6.5, 6.0])

    heading(doc, "目标业务场景")
    paragraph(doc, "本方案聚焦学习支持请求的接收、核对、分派、处理和复盘。一个请求只有在授权状态、服务规则版本和具名处理责任人均明确时，才进入处理队列；其余情况进入学习服务主任的人工复核队列。")
    table(doc, ["角色", "主要关注点", "一期责任"], [
        ["学习服务主任", "服务价值与试点边界", "批准范围、复核异常、决定继续或停止"],
        ["助教", "具体学习事件与现场习惯", "提出并补充请求"],
        ["课程运营员", "材料、交接和规则执行", "核对流程状态"],
        ["课程讲师", "学习内容与下游反馈", "处理具名任务"],
        ["教育数据专员", "权限、留存和审计", "审查数据控制"],
        ["平台工程师", "系统、版本和故障处理", "实现接口与日志"],
    ], [3.2, 6.5, 6.0])

    heading(doc, "转型方案与能力设计")
    paragraph(doc, "转型后以任务为协作单元。请求进入后先标准化字段，再执行授权、规则版本与责任人校验。控制通过的请求由系统分派给具名角色；控制不通过的请求不自动继续，而是携带原因进入人工复核。每一个决定都形成审计事件，为后续复盘提供依据。")
    table(doc, ["能力", "一期实现", "控制要求"], [
        ["请求结构化", "请求编号、区域、主题、来源角色", "不收集非必要个人信息"],
        ["规则控制", "必填规则版本并与当前版本比对", "版本不匹配必须阻断"],
        ["协作分派", "指定具名处理责任人", "无责任人必须人工复核"],
        ["审计复盘", "记录决定、原因、下一责任人与事件", "审计范围和留存期由数据负责人确认"],
    ], [3.2, 6.5, 6.0])

    heading(doc, "价值衡量与试点建议")
    paragraph(doc, "建议在试点前冻结指标口径，并按周抽取同类脱敏样本。关注首次响应时长、平均处理时长、返工次数、规则版本正确率、采用率和人工复核分钟数。指标变化应与对照组或试点前基线比较，避免将季节性波动或样本差异误认为方案收益。")
    table(doc, ["指标", "定义", "试点判定"], [
        ["处理时长", "从请求登记到具名责任人完成的时间", "不高于基线且异常可解释"],
        ["返工次数", "因缺信息、版本错误或责任不清导致的重复处理", "较基线下降"],
        ["版本正确率", "使用正确服务规则版本的任务占比", "关键场景达到 100%"],
        ["人工复核成本", "复核队列的人工分钟数", "转人工原因可分类并逐步优化"],
    ], [3.2, 7.2, 5.3])

    heading(doc, "实施路线与决策条件")
    table(doc, ["阶段", "周期", "关键产出", "继续条件"], [
        ["业务挖掘和调研", "第 1 至 2 周", "访谈、样例、指标口径、现状流程", "关键角色确认问题和范围"],
        ["需求与方案", "第 3 至 4 周", "需求矩阵、权限设计、试点计划", "验收条件和数据访问获批"],
        ["原型与测试", "第 5 至 7 周", "可运行原型、测试报告、异常样例", "控制逻辑通过评审"],
        ["受控试点", "第 8 至 10 周", "试点复盘、收益核对、交接建议", "指标、风险和退出条件满足"],
    ], [3.2, 2.5, 7.0, 3.0])
    paragraph(doc, "建议由学习服务主任主持阶段评审。若授权、规则版本、数据访问或关键责任人任一未确认，应暂停扩大范围并回到对应阶段补齐证据。")

    doc.core_properties.title = "学习支持与辅导协作行业分析报告"
    doc.core_properties.author = "Learning Support Collaboration Project"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
